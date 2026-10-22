import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from typing import List, Dict, Any, Optional
from datetime import datetime
import io
import logging

logger = logging.getLogger(__name__)


class ExcelExporter:
    def __init__(self):
        self.header_font = Font(bold=True, size=11)
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font_white = Font(bold=True, size=11, color="FFFFFF")
        self.money_format = '#,##0.00'
        self.date_format = 'YYYY-MM-DD'
        self.warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        self.error_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _sanitize_cell_value(self, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, str):
            if value.startswith(('=', '+', '-', '@')):
                return "'" + value
        return value

    def _apply_header_style(self, ws, row: int, max_col: int):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = self.header_font_white
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = self.thin_border

    def _apply_cell_style(self, ws, row: int, max_col: int, highlight_cols: set = None):
        highlight_cols = highlight_cols or set()
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = self.thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if col in highlight_cols:
                cell.fill = self.warning_fill

    def _auto_width(self, ws, max_col: int, min_width: int = 10, max_width: int = 50):
        for col in range(1, max_col + 1):
            max_length = min_width
            for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
                for cell in row:
                    if cell.value:
                        length = len(str(cell.value))
                        max_length = max(max_length, min(length, max_width))
            ws.column_dimensions[get_column_letter(col)].width = max_length + 2

    def export_receipts(self, receipts: List[Dict[str, Any]], 
                        include_ocr: bool = False) -> bytes:
        wb = openpyxl.Workbook()
        
        ws_receipts = wb.active
        ws_receipts.title = "Receipts"
        self._write_receipts_sheet(ws_receipts, receipts)
        
        ws_items = wb.create_sheet("Receipt Items")
        self._write_items_sheet(ws_items, receipts)
        
        ws_validation = wb.create_sheet("Validation")
        self._write_validation_sheet(ws_validation, receipts)
        
        if include_ocr:
            ws_ocr = wb.create_sheet("OCR Text")
            self._write_ocr_sheet(ws_ocr, receipts)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def export_single_receipt(self, receipt: Dict[str, Any]) -> bytes:
        wb = openpyxl.Workbook()
        
        ws = wb.active
        ws.title = "Receipt"
        self._write_single_receipt(ws, receipt)
        
        if receipt.get('items'):
            ws_items = wb.create_sheet("Items")
            self._write_single_receipt_items(ws_items, receipt)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def _write_receipts_sheet(self, ws, receipts: List[Dict[str, Any]]):
        headers = [
            'ID', 'Merchant Name', 'Branch', 'Address', 'City', 'Country',
            'Contact Number', 'Website/Email', 'Merchant TIN', 'Receipt Number',
            'Transaction Ref', 'Invoice Number', 'Order Number', 'Terminal Number',
            'Transaction Date', 'Transaction Time', 'Currency', 'Currency Symbol',
            'Payment Method', 'Card Brand', 'Masked Card Number', 'Transaction Status',
            'Cashier Name', 'Subtotal', 'Item Discounts', 'Transaction Discounts',
            'Total Discounts', 'Amount Before Tax', 'VAT/Tax', 'Taxable Sales',
            'VAT-Exempt Sales', 'Zero-Rated Sales', 'Other Taxes/Fees',
            'Service Charge', 'Delivery Fee', 'Packaging Fee', 'Tips/Gratuity',
            'Rounding Adjustment', 'Grand Total', 'Amount Paid', 'Cash Received',
            'Change Given', 'Outstanding Balance', 'Validation Status', 'Review Status',
            'Created At', 'Updated At', 'Image Path'
        ]
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(receipts) + 1}"
        
        money_cols = {23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40}
        
        for row_idx, receipt in enumerate(receipts, 2):
            row_data = [
                receipt.get('id'),
                receipt.get('merchant_name'),
                receipt.get('branch_name'),
                receipt.get('address'),
                receipt.get('city'),
                receipt.get('country'),
                receipt.get('contact_number'),
                receipt.get('website_email'),
                receipt.get('merchant_tin'),
                receipt.get('receipt_number'),
                receipt.get('transaction_ref'),
                receipt.get('invoice_number'),
                receipt.get('order_number'),
                receipt.get('terminal_number'),
                receipt.get('transaction_date'),
                receipt.get('transaction_time'),
                receipt.get('currency'),
                receipt.get('currency_symbol'),
                receipt.get('payment_method'),
                receipt.get('card_brand'),
                receipt.get('masked_card_number'),
                receipt.get('transaction_status'),
                receipt.get('cashier_name'),
                receipt.get('subtotal'),
                receipt.get('item_discounts'),
                receipt.get('transaction_discounts'),
                receipt.get('total_discounts'),
                receipt.get('amount_before_tax'),
                receipt.get('vat_tax'),
                receipt.get('taxable_sales'),
                receipt.get('vat_exempt_sales'),
                receipt.get('zero_rated_sales'),
                receipt.get('other_taxes_fees'),
                receipt.get('service_charge'),
                receipt.get('delivery_fee'),
                receipt.get('packaging_fee'),
                receipt.get('tips_gratuity'),
                receipt.get('rounding_adjustment'),
                receipt.get('grand_total'),
                receipt.get('amount_paid'),
                receipt.get('cash_received'),
                receipt.get('change_given'),
                receipt.get('outstanding_balance'),
                receipt.get('validation_status'),
                receipt.get('review_status'),
                receipt.get('created_at'),
                receipt.get('updated_at'),
                receipt.get('original_image_path'),
            ]
            
            highlight = set()
            if receipt.get('validation_status') in ('Needs Review', 'Incomplete'):
                highlight = set(range(1, len(headers) + 1))
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=self._sanitize_cell_value(value))
                if col in money_cols and isinstance(value, (int, float)):
                    cell.number_format = self.money_format
                elif col in (15, 16, 46, 47) and value:
                    cell.number_format = self.date_format
            
            self._apply_cell_style(ws, row_idx, len(headers), highlight)
        
        self._auto_width(ws, len(headers))

    def _write_items_sheet(self, ws, receipts: List[Dict[str, Any]]):
        headers = [
            'Receipt ID', 'Item ID', 'Description', 'Quantity', 'Unit Price',
            'Discount', 'Tax', 'Line Total', 'Extraction Evidence'
        ]
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        ws.freeze_panes = 'A2'
        
        money_cols = {5, 6, 7, 8}
        
        row_idx = 2
        for receipt in receipts:
            for item in receipt.get('items', []):
                row_data = [
                    receipt.get('id'),
                    item.get('id'),
                    item.get('description'),
                    item.get('quantity'),
                    item.get('unit_price'),
                    item.get('discount'),
                    item.get('tax'),
                    item.get('line_total'),
                    item.get('extraction_evidence'),
                ]
                
                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col, value=self._sanitize_cell_value(value))
                    if col in money_cols and isinstance(value, (int, float)):
                        cell.number_format = self.money_format
                
                self._apply_cell_style(ws, row_idx, len(headers))
                row_idx += 1
        
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_idx}"
        self._auto_width(ws, len(headers))

    def _write_validation_sheet(self, ws, receipts: List[Dict[str, Any]]):
        headers = [
            'Receipt ID', 'Merchant', 'Date', 'Total', 'Currency',
            'Validation Status', 'Warnings', 'Missing Fields', 'Checks Passed'
        ]
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        ws.freeze_panes = 'A2'
        
        for row_idx, receipt in enumerate(receipts, 2):
            warnings = receipt.get('warnings', [])
            if isinstance(warnings, str):
                warnings = [warnings]
            
            missing = []
            critical_fields = ['merchant_name', 'transaction_date', 'currency', 'grand_total']
            for field in critical_fields:
                if not receipt.get(field):
                    missing.append(field)
            
            checks = receipt.get('validation_checks', {})
            passed = sum(1 for v in checks.values() if v)
            total_checks = len(checks)
            
            row_data = [
                receipt.get('id'),
                receipt.get('merchant_name'),
                receipt.get('transaction_date'),
                receipt.get('grand_total'),
                receipt.get('currency'),
                receipt.get('validation_status'),
                '; '.join(warnings) if warnings else 'None',
                ', '.join(missing) if missing else 'None',
                f"{passed}/{total_checks}"
            ]
            
            highlight = set()
            if receipt.get('validation_status') in ('Needs Review', 'Incomplete'):
                highlight = {6, 7, 8}
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=self._sanitize_cell_value(value))
                if col == 4 and isinstance(value, (int, float)):
                    cell.number_format = self.money_format
            
            self._apply_cell_style(ws, row_idx, len(headers), highlight)
        
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(receipts) + 1}"
        self._auto_width(ws, len(headers))

    def _write_ocr_sheet(self, ws, receipts: List[Dict[str, Any]]):
        headers = ['Receipt ID', 'Merchant', 'Date', 'Raw OCR Text']
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        ws.freeze_panes = 'A2'
        
        for row_idx, receipt in enumerate(receipts, 2):
            row_data = [
                receipt.get('id'),
                receipt.get('merchant_name'),
                receipt.get('transaction_date'),
                receipt.get('raw_ocr_text', '')[:32000]
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=self._sanitize_cell_value(value))
                if col == 4:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            
            self._apply_cell_style(ws, row_idx, len(headers))
        
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(receipts) + 1}"
        self._auto_width(ws, len(headers))

    def _write_single_receipt(self, ws, receipt: Dict[str, Any]):
        headers = [
            'Field', 'Value', 'Currency'
        ]
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        
        fields = [
            ('Merchant Name', 'merchant_name', None),
            ('Branch', 'branch_name', None),
            ('Address', 'address', None),
            ('City', 'city', None),
            ('Country', 'country', None),
            ('Contact Number', 'contact_number', None),
            ('Website/Email', 'website_email', None),
            ('Merchant TIN', 'merchant_tin', None),
            ('Receipt Number', 'receipt_number', None),
            ('Transaction Ref', 'transaction_ref', None),
            ('Invoice Number', 'invoice_number', None),
            ('Order Number', 'order_number', None),
            ('Terminal Number', 'terminal_number', None),
            ('Transaction Date', 'transaction_date', None),
            ('Transaction Time', 'transaction_time', None),
            ('Currency', 'currency', None),
            ('Currency Symbol', 'currency_symbol', None),
            ('Payment Method', 'payment_method', None),
            ('Card Brand', 'card_brand', None),
            ('Masked Card Number', 'masked_card_number', None),
            ('Transaction Status', 'transaction_status', None),
            ('Cashier Name', 'cashier_name', None),
            ('Subtotal', 'subtotal', 'currency'),
            ('Item Discounts', 'item_discounts', 'currency'),
            ('Transaction Discounts', 'transaction_discounts', 'currency'),
            ('Total Discounts', 'total_discounts', 'currency'),
            ('Amount Before Tax', 'amount_before_tax', 'currency'),
            ('VAT/Tax', 'vat_tax', 'currency'),
            ('Taxable Sales', 'taxable_sales', 'currency'),
            ('VAT-Exempt Sales', 'vat_exempt_sales', 'currency'),
            ('Zero-Rated Sales', 'zero_rated_sales', 'currency'),
            ('Other Taxes/Fees', 'other_taxes_fees', 'currency'),
            ('Service Charge', 'service_charge', 'currency'),
            ('Delivery Fee', 'delivery_fee', 'currency'),
            ('Packaging Fee', 'packaging_fee', 'currency'),
            ('Tips/Gratuity', 'tips_gratuity', 'currency'),
            ('Rounding Adjustment', 'rounding_adjustment', 'currency'),
            ('Grand Total', 'grand_total', 'currency'),
            ('Amount Paid', 'amount_paid', 'currency'),
            ('Cash Received', 'cash_received', 'currency'),
            ('Change Given', 'change_given', 'currency'),
            ('Outstanding Balance', 'outstanding_balance', 'currency'),
            ('Validation Status', 'validation_status', None),
            ('Review Status', 'review_status', None),
            ('Created At', 'created_at', None),
            ('Updated At', 'updated_at', None),
        ]
        
        currency = receipt.get('currency', 'PHP')
        
        for row_idx, (label, key, curr_key) in enumerate(fields, 2):
            value = receipt.get(key)
            curr = receipt.get(curr_key) if curr_key else currency
            
            ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row_idx, column=1).border = self.thin_border
            
            cell = ws.cell(row=row_idx, column=2, value=self._sanitize_cell_value(value))
            cell.border = self.thin_border
            if isinstance(value, (int, float)):
                cell.number_format = self.money_format
            
            ws.cell(row=row_idx, column=3, value=curr).border = self.thin_border
        
        self._auto_width(ws, 3)

    def _write_single_receipt_items(self, ws, receipt: Dict[str, Any]):
        headers = [
            'Description', 'Quantity', 'Unit Price', 'Discount', 'Tax', 'Line Total', 'Extraction Evidence'
        ]
        
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        self._apply_header_style(ws, 1, len(headers))
        ws.freeze_panes = 'A2'
        
        money_cols = {3, 4, 5, 6}
        
        for row_idx, item in enumerate(receipt.get('items', []), 2):
            row_data = [
                item.get('description'),
                item.get('quantity'),
                item.get('unit_price'),
                item.get('discount'),
                item.get('tax'),
                item.get('line_total'),
                item.get('extraction_evidence'),
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=self._sanitize_cell_value(value))
                cell.border = self.thin_border
                if col in money_cols and isinstance(value, (int, float)):
                    cell.number_format = self.money_format
        
        self._auto_width(ws, len(headers))


def get_exporter() -> ExcelExporter:
    return ExcelExporter()